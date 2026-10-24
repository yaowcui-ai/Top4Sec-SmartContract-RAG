"""Build the operator manual with python-docx (use the bundled document runtime)."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]

def font(style,name='Microsoft YaHei',size=11):
    style.font.name=name;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')

def text(paragraph,value):
    for part in re.split(r'(`[^`]+`)',value):
        run=paragraph.add_run(part.strip('`') if part.startswith('`') else part)
        if part.startswith('`'):run.font.name='Consolas';run.font.size=Pt(10)

def build():
    doc=Document();s=doc.sections[0]
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.7);s.bottom_margin=Inches(.65);s.left_margin=Inches(.75);s.right_margin=Inches(.75)
    for name,size in [('Normal',11),('Title',21),('Heading 1',15),('Heading 2',12)]:font(doc.styles[name],size=size)
    normal=doc.styles['Normal'];normal.paragraph_format.line_spacing=1.18;normal.paragraph_format.space_after=Pt(6)
    for name in ['Title','Heading 1','Heading 2']:
        pf=doc.styles[name].paragraph_format;pf.keep_with_next=True;pf.space_before=Pt(12);pf.space_after=Pt(7)
    code=doc.styles.add_style('Command',1);font(code,'Consolas',9.3)
    code.paragraph_format.space_after=Pt(3);code.paragraph_format.line_spacing=1.05
    code.paragraph_format.left_indent=Inches(.12);code.paragraph_format.right_indent=Inches(.08)
    header=s.header.paragraphs[0];header.text='四大顶会区块链智能合约论文问答系统  操作手册';font(header.style,size=9)
    footer=s.footer.paragraphs[0];footer.alignment=2
    footer.add_run('第 ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页')
    lines=(ROOT/'docs/操作手册.md').read_text(encoding='utf-8').splitlines();i=0;in_code=False
    while i<len(lines):
        line=lines[i];i+=1
        if line.startswith('```'):in_code=not in_code;continue
        if in_code:
            p=doc.add_paragraph(line,style='Command');p.paragraph_format.keep_together=True
            p.paragraph_format.keep_with_next=i<len(lines) and not lines[i].startswith('```')
            continue
        if not line.strip():continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
            rows=[r for r in rows if not re.fullmatch(r'[| :\-]+',r)]
            matrix=[[c.strip().replace('`','') for c in r.strip().strip('|').split('|')] for r in rows]
            t=doc.add_table(rows=0,cols=len(matrix[0]));t.autofit=False
            widths=[2.35,4.65] if len(matrix[0])==2 else [1.55,2.05,3.4]
            for column,width in zip(t.columns,widths):column.width=Inches(width)
            for row_index,values in enumerate(matrix):
                cells=t.add_row().cells
                for j,value in enumerate(values):
                    cells[j].width=Inches(widths[j]);cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cells[j].paragraphs[0];p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
                    r=p.add_run(value);r.font.size=Pt(10);r.bold=row_index==0
                    pr=cells[j]._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
                    for side in ['top','bottom','left','right']:
                        x=OxmlElement('w:'+side);x.set(qn('w:w'),'90');x.set(qn('w:type'),'dxa');m.append(x)
                    pr.append(m)
                    if row_index==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');pr.append(sh)
                trpr=t.rows[-1]._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
                if row_index==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
        elif line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
        elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
        elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
        else:text(doc.add_paragraph(),line)
    out=ROOT/'docs/四大顶会区块链智能合约论文问答系统操作手册.docx';doc.save(out);print(out)

if __name__=='__main__':build()
